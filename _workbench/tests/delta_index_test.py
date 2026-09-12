import hashlib, json, os, sys, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import index_doc as idx
from render_document import render
from office_convert import office_bin
from verify_cards import verify
from merge_cards import merge
from PIL import Image
from openpyxl import Workbook
import pymupdf

class DeltaIndex(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
  (self.root/'_analysis/index').mkdir(parents=True);(self.root/'_analysis/cards').mkdir()
  self.env=patch.dict(os.environ,{'QLIST_INDEX_AI':'0'});self.env.start();self.addCleanup(self.env.stop)
 def save_index(self,file):
  body=idx.build(str(file),'未分輪');self.write_index(file.name,body);return body
 def write_index(self,name,body):
  (self.root/'_analysis/index'/(name+'.index.json')).write_text(json.dumps(body))
 def test_mixed_pdf_tables_image_union_and_native_only(self):
  pic=self.root/'photo.jpg';Image.new('RGB',(100,100),'navy').save(pic)
  file=self.root/'mixed.pdf';d=pymupdf.open();p=d.new_page(width=500,height=500)
  p.insert_text((20,30),'Native explanatory text longer than thirty characters.')
  p.insert_image(pymupdf.Rect(0,100,300,400),filename=str(pic));p.insert_image(pymupdf.Rect(0,100,300,400),filename=str(pic))
  for x in [320,400,480]:p.draw_line((x,100),(x,160))
  for y in [100,130,160]:p.draw_line((320,y),(480,y))
  for x,y,t in [(325,120,'Item'),(405,120,'Value'),(325,150,'Revenue'),(405,150,'275')]:p.insert_text((x,y),t,fontsize=10)
  d.new_page().insert_text((20,30),'Text-only source page with more than thirty characters.');d.save(file);d.close()
  with patch.object(idx,'describe',side_effect=AssertionError('native index called model')):body=self.save_index(file)
  p=body['pages'][0];self.assertEqual(p['images'],2);self.assertAlmostEqual(p['image_area_ratio'],.36);self.assertTrue(p['needs_visual']);self.assertFalse(p['needs_ocr']);self.assertIn('275',str(p['tables']))
  result=render(self.root,file.name);self.assertEqual([r['page'] for r in result],[1]);self.assertTrue((self.root/result[0]['path']).is_file())
  with self.assertRaises(ValueError):render(self.root,file.name,2)
  file.write_bytes(file.read_bytes()+b'\n')
  with self.assertRaisesRegex(ValueError,'重建索引'):render(self.root,file.name,1)
  with self.assertRaises(ValueError):render(self.root,'../mixed.pdf',1)
 def test_images_render_with_no_fabricated_text(self):
  from pillow_heif import register_heif_opener
  register_heif_opener()
  for ext in ['jpg','png','webp','heic']:
   f=self.root/('photo.'+ext);im=Image.new('RGB',(80,50),'red');im.save(f,format='HEIF' if ext=='heic' else None)
   body=self.save_index(f);self.assertEqual(body['pages'][0]['text'],'');self.assertEqual(body['needs_visual_pages'],1)
   result=render(self.root,f.name,1)
   with Image.open(self.root/result[0]['path']) as rendered:self.assertEqual(rendered.size,(80,50))
   self.assertEqual(render(self.root,f.name,1),result)
 def test_nested_render_source_and_duplicate_guard(self):
  nested=self.root/'round1/attachments';nested.mkdir(parents=True);f=nested/'nested.jpg';Image.new('RGB',(80,50),'red').save(f)
  body=idx.build(str(f),'R1');self.write_index(f.name,body);self.assertEqual(len(render(self.root,f.name)),1)
  (self.root/'round1/nested.jpg').write_bytes(f.read_bytes())
  with self.assertRaisesRegex(ValueError,'同名'):render(self.root,f.name)
 def make_pptx(self):
  from pptx import Presentation
  from pptx.chart.data import CategoryChartData
  from pptx.enum.chart import XL_CHART_TYPE
  from pptx.util import Inches
  photo=self.root/'chart-photo.jpg';Image.new('RGB',(80,50),'blue').save(photo)
  prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);data=CategoryChartData();data.categories=['2024','2025'];data.add_series('Revenue',[125,275]);slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(1),Inches(5),Inches(3),data);slide.shapes.add_picture(str(photo),Inches(7),Inches(1));file=self.root/'deck.pptx';prs.save(file);return file
 def test_pptx_native_chart_and_missing_converter(self):
  file=self.make_pptx();body=self.save_index(file);page=body['pages'][0];self.assertTrue(page['has_pictures']);self.assertTrue(page['needs_visual']);self.assertIn('275',page['text']);self.assertEqual(page['charts'][0]['categories'],['2024','2025'])
  with patch('office_convert.office_bin',return_value=None):
   with self.assertRaisesRegex(RuntimeError,'LibreOffice'):render(self.root,file.name)
 @unittest.skipUnless(office_bin(),'optional LibreOffice not installed')
 def test_real_pptx_conversion_and_formula_copy(self):
  file=self.make_pptx();self.save_index(file);before=hashlib.sha256(file.read_bytes()).digest();result=render(self.root,file.name);self.assertEqual(len(result),1);self.assertEqual(hashlib.sha256(file.read_bytes()).digest(),before)
  file=self.make_xlsx();before=file.read_bytes();body=idx.index_xlsx(file);c=body['sheets'][0]['cells'][-1];self.assertEqual(c['value'],210);self.assertEqual(c['formula'],'=B1*2.1');self.assertEqual(body['formula_cache']['missing'],0);self.assertEqual(file.read_bytes(),before)
 def make_xlsx(self):
  file=self.root/'model.xlsx';w=Workbook();w.active.append(['Revenue',100]);w.active.append(['Forecast','=B1*2.1']);w.save(file);return file
 def test_formula_fallback_retains_formula(self):
  file=self.make_xlsx()
  with patch('office_convert.office_bin',return_value=None):body=idx.index_xlsx(file)
  self.assertEqual(body['formula_cache']['missing'],1);c=body['sheets'][0]['cells'][-1];self.assertIsNone(c['value']);self.assertEqual(c['formula'],'=B1*2.1');self.assertTrue(c['cache_missing']);self.assertIn('值未快取',body['sheets'][0]['text'])
 def test_segment_coordinates(self):
  f=self.root/'long.txt';f.write_text('a'*5001);body=self.save_index(f);self.assertEqual(body['page_count'],2);self.assertEqual(body['pages'][1]['loc_kind'],'segment')
 def test_statements_exact_quotes_and_merge_retention(self):
  self.write_index('claims.pdf',{'kind':'pdf','page_count':2,'pages':[{'n':1,'text':'Revenue 275. No pending litigation. We employ 42 staff.'},{'n':2,'text':'The founder is full time.'}]})
  cards=self.root/'_analysis/cards'
  for n,quote in [(1,'No pending litigation.'),(2,'The founder is full time.')]:
   (cards/f'claims.pdf.part-0{n}.md').write_text(f'## 摘要\nClaim {n}.\n## 關鍵數字表\n| 項目 | 數值 | 出處 |\n'+('| Revenue | 275 | p.1 |\n' if n==1 else '')+f'## 重要陳述（非數字）\n| 陳述 | 原文摘錄 | 出處 |\n| --- | --- | --- |\n| 承諾 | {quote} | p.{n} |\n## 未明名詞與疑點\n## 覆蓋聲明\n<!-- coverage: '+json.dumps({'pages':[n,n],'sheets':[],'unreadable':[]})+' -->')
  merge(self.root,'claims.pdf');out=cards/'claims.pdf.md';text=out.read_text();self.assertIn('No pending litigation.',text);self.assertIn('The founder is full time.',text)
  report=verify(self.root)[0];self.assertEqual(report['statement_rows'],2);self.assertEqual(report['hit_rate'],1)
  out.write_text(text.replace('No pending litigation.','No outstanding litigation.').replace('275 |','276 |'));report=verify(self.root)[0];self.assertEqual(len(report['miss']),2)
  out.write_text(text.replace('No pending litigation.','No  pending\t litigation.'));self.assertEqual(verify(self.root)[0]['hit_rate'],1)
  out.write_text(text.replace('No pending litigation.','We have 42 staff.'));self.assertEqual(len(verify(self.root)[0]['miss']),1)
  out.write_text(text.replace('No pending litigation.','N'*61));self.assertEqual(len(verify(self.root)[0]['miss']),1)
if __name__=='__main__':unittest.main()

import sys, json
from openpyxl import load_workbook

# 讀 xlsx 供工作台預覽用。路徑走 argv[1]，輸出 JSON 到 stdout。
# 上限：每表 80 列 x 30 欄、最多 10 張工作表，避免大檔塞爆前端。
MAXR, MAXC, MAXSHEETS = 80, 30, 10

def cell(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)

def main():
    wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets[:MAXSHEETS]:
        raw, truncated = [], False
        for ri, row in enumerate(ws.iter_rows(values_only=True)):
            if ri >= MAXR:
                truncated = True
                break
            raw.append([cell(v) for v in row[:MAXC]])
        while raw and not any(c.strip() for c in raw[-1]):
            raw.pop()
        # 只有真內容填滿到列上限才算截斷；砍掉尾端空列後不足上限＝沒截斷
        truncated = truncated and len(raw) >= MAXR
        width = 0
        for r in raw:
            for ci in range(len(r) - 1, -1, -1):
                if r[ci].strip():
                    width = max(width, ci + 1)
                    break
        rows = [(r + [''] * width)[:width] for r in raw]
        sheets.append({'name': ws.title, 'rows': rows, 'truncated': truncated})
    wb.close()
    sys.stdout.write(json.dumps({'sheets': sheets}, ensure_ascii=False))

main()

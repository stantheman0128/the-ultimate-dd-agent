#!/usr/bin/env python3
"""Data Room 文件索引：上傳時抽取「每頁文字 / 每格值與公式」，寫到 <案件>/_analysis/index/<檔名>.index.json。

用法：
  python3 _workbench/index_doc.py <案件資料夾>                 # 索引全部文件（已索引且未變動者略過）
  python3 _workbench/index_doc.py <案件資料夾> --file <路徑>   # 只索引這一份（相對案件資料夾或絕對路徑）
  python3 _workbench/index_doc.py <案件資料夾> --force         # 全部重建

索引是引擎與工作台共用的「可定位層」：
- card-extractor 讀它做字卡（掃描頁 needs_ocr=true 時才回頭用視覺讀原檔）
- /api/search 用它搜原文、/api/page 用它回單頁、/api/ask 用它組 context 並標出處 [檔名 p.N] / [檔名 工作表!B4]
"""
import json, os, sys, time, datetime, re

SKIP_PREFIX = ('.', '~$')
DOC_EXT = {'.pdf', '.xlsx', '.xlsm', '.pptx', '.docx', '.txt', '.md', '.csv'}
MAX_CELLS_PER_SHEET = 40000
OCR_MIN_CHARS = 30  # 一頁文字少於此且有圖 → 判為掃描頁

def is_qlist(name):
    return re.search(r'q-?list', name, re.I) is not None

def list_docs(deal):
    out = []
    for e in sorted(os.listdir(deal)):
        p = os.path.join(deal, e)
        if re.fullmatch(r'round\d+', e) and os.path.isdir(p):
            for f in sorted(os.listdir(p)):
                if not f.startswith(SKIP_PREFIX):
                    out.append((os.path.join(p, f), 'R' + e[5:]))
        elif os.path.isfile(p) and not e.startswith(SKIP_PREFIX) and e != '_notes.md' and not is_qlist(e):
            out.append((p, '未分輪'))
    return [(p, r) for p, r in out if os.path.splitext(p)[1].lower() in DOC_EXT]

def index_pdf(path):
    import pymupdf
    doc = pymupdf.open(path)
    pages, ocr_pages = [], 0
    has_tess = bool(os.environ.get('TESSDATA_PREFIX')) or os.path.exists('/usr/share/tesseract-ocr')
    for i, page in enumerate(doc):
        text = page.get_text('text') or ''
        needs_ocr = len(text.strip()) < OCR_MIN_CHARS and bool(page.get_images(full=False))
        if needs_ocr and has_tess:
            try:
                tp = page.get_textpage_ocr(language='chi_tra+eng', dpi=200, full=True)
                text = page.get_text('text', textpage=tp) or text
                needs_ocr = len(text.strip()) < OCR_MIN_CHARS
            except Exception:
                pass
        if needs_ocr:
            ocr_pages += 1
        pages.append({'n': i + 1, 'text': text, 'chars': len(text), 'needs_ocr': needs_ocr})
    return {'kind': 'pdf', 'pages': pages, 'page_count': len(pages), 'needs_ocr_pages': ocr_pages}

def index_xlsx(path):
    from openpyxl import load_workbook
    wb_v = load_workbook(path, read_only=True, data_only=True)
    wb_f = load_workbook(path, read_only=False, data_only=False)
    sheets = []
    for ws in wb_f.worksheets:
        ws_v = wb_v[ws.title] if ws.title in wb_v.sheetnames else None
        vals = {}
        if ws_v is not None:
            for row in ws_v.iter_rows():
                for c in row:
                    if c.value is not None:
                        vals[c.coordinate] = c.value
        cells, truncated = [], False
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                formula = v if isinstance(v, str) and v.startswith('=') else None
                value = vals.get(c.coordinate, None if formula else v)
                if isinstance(value, (datetime.date, datetime.datetime)):
                    value = value.isoformat()
                cells.append({'ref': c.coordinate, 'value': value, 'formula': formula})
                if len(cells) >= MAX_CELLS_PER_SHEET:
                    truncated = True
                    break
            if truncated:
                break
        text = '\n'.join(f"{c['ref']}\t{c['value']}" + (f"\t{c['formula']}" if c['formula'] else '') for c in cells)
        sheets.append({'name': ws.title, 'cells': cells, 'cell_count': len(cells), 'truncated': truncated,
                       'dims': ws.dimensions, 'text': text})
    wb_v.close()
    return {'kind': 'xlsx', 'sheets': sheets, 'sheet_count': len(sheets)}

def index_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for i, s in enumerate(prs.slides):
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows:
                    parts.append('\t'.join(c.text for c in r.cells))
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            nt = s.notes_slide.notes_text_frame.text.strip()
            if nt:
                parts.append('[notes] ' + nt)
        text = '\n'.join(parts)
        slides.append({'n': i + 1, 'text': text, 'chars': len(text), 'needs_ocr': len(text.strip()) < OCR_MIN_CHARS})
    return {'kind': 'pptx', 'pages': slides, 'page_count': len(slides)}

def index_docx(path):
    try:
        import docx
    except ImportError:
        return {'kind': 'docx', 'pages': [], 'page_count': 0, 'error': 'python-docx not installed'}
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            parts.append('\t'.join(c.text for c in r.cells))
    text = '\n'.join(parts)
    # docx 沒有穩定頁碼：以 3,000 字切「頁」
    chunks = [text[i:i + 3000] for i in range(0, max(len(text), 1), 3000)]
    return {'kind': 'docx', 'pages': [{'n': i + 1, 'text': c, 'chars': len(c), 'needs_ocr': False} for i, c in enumerate(chunks)],
            'page_count': len(chunks), 'note': '頁碼為 3000 字切分，非原始分頁'}

def index_text(path):
    with open(path, encoding='utf-8', errors='replace') as f:
        text = f.read()
    chunks = [text[i:i + 3000] for i in range(0, max(len(text), 1), 3000)]
    return {'kind': 'text', 'pages': [{'n': i + 1, 'text': c, 'chars': len(c), 'needs_ocr': False} for i, c in enumerate(chunks)],
            'page_count': len(chunks)}

def build(path, round_label):
    ext = os.path.splitext(path)[1].lower()
    t0 = time.time()
    if ext == '.pdf':
        body = index_pdf(path)
    elif ext in ('.xlsx', '.xlsm'):
        body = index_xlsx(path)
    elif ext == '.pptx':
        body = index_pptx(path)
    elif ext == '.docx':
        body = index_docx(path)
    else:
        body = index_text(path)
    st = os.stat(path)
    body.update({'file': os.path.basename(path), 'round': round_label, 'size': st.st_size, 'mtime': st.st_mtime,
                 'indexed_at': datetime.datetime.now().isoformat(timespec='seconds'), 'seconds': round(time.time() - t0, 2)})
    body['total_chars'] = sum(p.get('chars', 0) for p in body.get('pages', [])) + sum(len(s.get('text', '')) for s in body.get('sheets', []))
    return body

def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not args:
        print(__doc__); sys.exit(2)
    deal = os.path.abspath(args[0])
    force = '--force' in sys.argv
    only = None
    if '--file' in sys.argv:
        only = sys.argv[sys.argv.index('--file') + 1]
        if not os.path.isabs(only) and not os.path.exists(only):
            only = os.path.join(deal, only)  # 相對案件資料夾的路徑（例：round1/x.pdf）
        only = os.path.abspath(only)
    out_dir = os.path.join(deal, '_analysis', 'index')
    os.makedirs(out_dir, exist_ok=True)
    results = []
    for path, rnd in list_docs(deal):
        if only and os.path.abspath(path) != only:
            continue
        out = os.path.join(out_dir, os.path.basename(path) + '.index.json')
        st = os.stat(path)
        if not force and os.path.exists(out):
            try:
                old = json.load(open(out, encoding='utf-8'))
                if old.get('mtime') == st.st_mtime and old.get('size') == st.st_size:
                    results.append({'file': os.path.basename(path), 'skipped': True, 'pages': old.get('page_count') or old.get('sheet_count')})
                    continue
            except Exception:
                pass
        try:
            body = build(path, rnd)
            with open(out, 'w', encoding='utf-8') as f:
                json.dump(body, f, ensure_ascii=False)
            results.append({'file': body['file'], 'kind': body['kind'], 'pages': body.get('page_count') or body.get('sheet_count'),
                            'needs_ocr_pages': body.get('needs_ocr_pages', 0), 'chars': body['total_chars'], 'seconds': body['seconds']})
        except Exception as e:  # noqa
            results.append({'file': os.path.basename(path), 'error': str(e)[:300]})
    print(json.dumps({'deal': os.path.basename(deal), 'results': results}, ensure_ascii=False))

if __name__ == '__main__':
    main()

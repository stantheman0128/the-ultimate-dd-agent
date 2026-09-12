#!/usr/bin/env python3
"""stdin JSON {path, rows, sheet} -> xlsx"""
import json, sys
import openpyxl
from openpyxl.styles import Font, Alignment

spec = json.load(sys.stdin)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = spec.get("sheet", "Sheet1")
for row in spec["rows"]:
    ws.append(row)
for cell in ws[1]:
    cell.font = Font(bold=True)
widths = [6, 14, 14, 80, 30, 8, 14, 50]
for i, w in enumerate(widths[: ws.max_column]):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(spec["path"])
print("ok")
